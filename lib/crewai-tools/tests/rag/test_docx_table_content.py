from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
import pytest
import requests

from crewai_tools.rag.loaders.docx_loader import DOCXLoader
from crewai_tools.rag.source_content import SourceContent


@pytest.mark.parametrize("source_kind", ["file", "url"])
@pytest.mark.parametrize("layout", ["table_only", "mixed", "nested"])
def test_docx_table_text_is_retained_in_document_order(
    source_kind: str, layout: str, tmp_path: Path
) -> None:
    """Retain table-only, mixed and nested content from local and URL sources."""
    document = Document()
    if layout != "table_only":
        document.add_paragraph("Before the table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Example item"
    table.cell(1, 1).text = "42 EUR"
    if layout == "nested":
        cell = table.cell(1, 0)
        nested = cell.add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Nested detail"
        cell.add_paragraph("After nested detail")
    if layout != "table_only":
        document.add_paragraph("After the table")
    data = BytesIO()
    document.save(data)
    path = tmp_path / "prices.docx"
    path.write_bytes(data.getvalue())
    source = str(path) if source_kind == "file" else "https://example.com/prices.docx"
    response = requests.Response()
    response.status_code = 200
    response._content = data.getvalue()

    with patch("crewai_tools.rag.loaders.docx_loader.safe_get", return_value=response):
        result = DOCXLoader().load(SourceContent(source))

    assert "Product | Price" in result.content
    assert "Example item" in result.content
    assert "42 EUR" in result.content
    assert result.metadata["tables"] == 1
    assert result.source == source
    if layout != "table_only":
        assert result.content.index("Before the table") < result.content.index(
            "Product"
        )
        assert result.content.index("42 EUR") < result.content.index("After the table")
    if layout == "nested":
        assert result.content.index("Example item") < result.content.index(
            "Nested detail"
        )
        assert result.content.index("Nested detail") < result.content.index(
            "After nested detail"
        )


@pytest.mark.parametrize("source_kind", ["file", "url"])
@pytest.mark.parametrize(
    ("merge_end", "expected_rows"),
    [
        (
            (0, 1),
            [
                "Merged value\nNested detail | Repeated",
                "Repeated | Repeated | Repeated",
            ],
        ),
        (
            (1, 0),
            [
                "Merged value\nNested detail | Repeated | Repeated",
                "Repeated | Repeated",
            ],
        ),
        ((1, 1), ["Merged value\nNested detail | Repeated", "Repeated"]),
    ],
    ids=["horizontal", "vertical", "rectangular"],
)
def test_docx_merged_cells_are_extracted_once(
    source_kind: str,
    merge_end: tuple[int, int],
    expected_rows: list[str],
    tmp_path: Path,
) -> None:
    """Extract each merged cell once without dropping equal text in distinct cells."""
    document = Document()
    document.add_paragraph("Before the table")
    table = document.add_table(rows=2, cols=3)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "Repeated"
    merged = table.cell(0, 0).merge(table.cell(*merge_end))
    merged.text = "Merged value"
    merged.add_table(rows=1, cols=1).cell(0, 0).text = "Nested detail"
    document.add_paragraph("After the table")
    data = BytesIO()
    document.save(data)
    path = tmp_path / "merged.docx"
    path.write_bytes(data.getvalue())
    source = str(path) if source_kind == "file" else "https://example.com/merged.docx"
    response = requests.Response()
    response.status_code = 200
    response._content = data.getvalue()
    loader = DOCXLoader()

    with patch("crewai_tools.rag.loaders.docx_loader.safe_get", return_value=response):
        result = loader.load(SourceContent(source))

    expected = "\n".join(["Before the table", *expected_rows, "After the table"])
    assert result.content == expected
    assert result.doc_id == loader.generate_doc_id(source_ref=source, content=expected)
    assert result.source == source
    assert result.metadata == {"format": "docx", "paragraphs": 2, "tables": 1}
