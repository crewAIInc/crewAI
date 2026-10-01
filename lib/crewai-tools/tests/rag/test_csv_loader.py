from datetime import datetime
import os
import sys
import tempfile
from io import BytesIO
from types import ModuleType
from unittest.mock import Mock, patch

from openpyxl import Workbook

from crewai_tools.rag.base_loader import LoaderResult
from crewai_tools.rag.data_types import DataType, DataTypes
from crewai_tools.rag.loaders.csv_loader import _EXCEL_URL_MAX_BYTES, CSVLoader
from crewai_tools.rag.source_content import SourceContent
import pytest


@pytest.fixture
def temp_csv_file():
    created_files = []

    def _create(content: str):
        f = tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False)
        f.write(content)
        f.close()
        created_files.append(f.name)
        return f.name

    yield _create

    for path in created_files:
        os.unlink(path)


class TestCSVLoader:
    @pytest.mark.parametrize("suffix", [".csv", ".xls", ".xlsx"])
    def test_detects_tabular_file_extensions(self, tmp_path, suffix):
        path = tmp_path / f"report{suffix}"
        path.touch()

        assert DataTypes.from_content(path) is DataType.CSV

    def test_load_csv_from_file(self, temp_csv_file):
        path = temp_csv_file("name,age,city\nJohn,25,New York\nJane,30,Chicago")
        loader = CSVLoader()
        result = loader.load(SourceContent(path))

        assert isinstance(result, LoaderResult)
        assert "Headers: name | age | city" in result.content
        assert "Row 1: name: John | age: 25 | city: New York" in result.content
        assert "Row 2: name: Jane | age: 30 | city: Chicago" in result.content
        assert result.metadata == {
            "format": "csv",
            "columns": ["name", "age", "city"],
            "rows": 2,
        }
        assert result.source == path
        assert result.doc_id

    def test_load_csv_with_empty_values(self, temp_csv_file):
        path = temp_csv_file("name,age,city\nJohn,,New York\n,30,")
        result = CSVLoader().load(SourceContent(path))

        assert "Row 1: name: John | city: New York" in result.content
        assert "Row 2: age: 30" in result.content
        assert result.metadata["rows"] == 2

    def test_load_csv_malformed(self, temp_csv_file):
        path = temp_csv_file('invalid,csv\nunclosed quote "missing')
        result = CSVLoader().load(SourceContent(path))

        assert "Headers: invalid | csv" in result.content
        assert 'Row 1: invalid: unclosed quote "missing' in result.content
        assert result.metadata["columns"] == ["invalid", "csv"]

    def test_load_csv_empty_file(self, temp_csv_file):
        path = temp_csv_file("")
        result = CSVLoader().load(SourceContent(path))

        assert result.content == ""
        assert result.metadata["rows"] == 0

    def test_load_csv_text_input(self):
        raw_csv = "col1,col2\nvalue1,value2\nvalue3,value4"
        result = CSVLoader().load(SourceContent(raw_csv))

        assert "Headers: col1 | col2" in result.content
        assert "Row 1: col1: value1 | col2: value2" in result.content
        assert "Row 2: col1: value3 | col2: value4" in result.content
        assert result.metadata["columns"] == ["col1", "col2"]
        assert result.metadata["rows"] == 2

    def test_doc_id_is_deterministic(self, temp_csv_file):
        path = temp_csv_file("name,value\ntest,123")
        loader = CSVLoader()

        result1 = loader.load(SourceContent(path))
        result2 = loader.load(SourceContent(path))

        assert result1.doc_id == result2.doc_id

    @patch("crewai_tools.security.safe_requests._raw_get")
    def test_load_csv_from_url(self, mock_get):
        mock_get.return_value = Mock(
            text="name,value\ntest,123", raise_for_status=Mock(return_value=None)
        )

        result = CSVLoader().load(SourceContent("https://example.com/data.csv"))

        assert "Headers: name | value" in result.content
        assert "Row 1: name: test | value: 123" in result.content
        headers = mock_get.call_args[1]["headers"]
        assert "text/csv" in headers["Accept"]
        assert "crewai-tools CSVLoader" in headers["User-Agent"]

    @patch("crewai_tools.security.safe_requests._raw_get")
    def test_load_csv_with_custom_headers(self, mock_get):
        mock_get.return_value = Mock(
            text="data,value\ntest,456", raise_for_status=Mock(return_value=None)
        )
        headers = {"Authorization": "Bearer token", "Custom-Header": "value"}
        result = CSVLoader().load(
            SourceContent("https://example.com/data.csv"), headers=headers
        )

        assert "Headers: data | value" in result.content
        assert mock_get.call_args[1]["headers"] == headers

    @patch("crewai_tools.security.safe_requests._raw_get")
    def test_csv_loader_handles_network_errors(self, mock_get):
        mock_get.side_effect = Exception("Network error")
        loader = CSVLoader()

        with pytest.raises(ValueError, match="Error fetching content from URL"):
            loader.load(SourceContent("https://example.com/data.csv"))

    @patch("crewai_tools.security.safe_requests._raw_get")
    def test_csv_loader_handles_http_error(self, mock_get):
        mock_get.return_value = Mock()
        mock_get.return_value.raise_for_status.side_effect = Exception("404 Not Found")
        loader = CSVLoader()

        with pytest.raises(ValueError, match="Error fetching content from URL"):
            loader.load(SourceContent("https://example.com/notfound.csv"))

    def test_load_xlsx_from_file(self, tmp_path):
        path = tmp_path / "report.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sales"
        worksheet.append(["name", "revenue"])
        worksheet.append(["North", 120])
        workbook.create_sheet("Empty")
        workbook.save(path)

        result = CSVLoader().load(SourceContent(path))

        assert "Sheet: Sales" in result.content
        assert "Headers: name | revenue" in result.content
        assert "Row 1: name: North | revenue: 120" in result.content
        assert result.metadata == {
            "format": "xlsx",
            "sheets": [
                {"name": "Sales", "columns": ["name", "revenue"], "rows": 1},
                {"name": "Empty", "columns": [], "rows": 0},
            ],
        }

    def test_load_xlsx_with_hash_in_local_name(self, tmp_path):
        path = tmp_path / "report#1.xlsx"
        workbook = Workbook()
        workbook.active.append(["name"])
        workbook.active.append(["Hashed"])
        workbook.save(path)

        result = CSVLoader().load(SourceContent(path))

        assert "Row 1: name: Hashed" in result.content
        assert result.metadata["format"] == "xlsx"

    def test_load_xls_uses_xlrd(self, monkeypatch):
        class Cell:
            def __init__(self, ctype: int, value: object) -> None:
                self.ctype = ctype
                self.value = value

        class Worksheet:
            name = "Legacy"
            nrows = 2

            @staticmethod
            def row(index: int) -> list[Cell]:
                return [
                    [Cell(1, "name"), Cell(1, "revenue"), Cell(1, "opened")],
                    [Cell(1, "South"), Cell(2, 80.0), Cell(3, 44927.0)],
                ][index]

        class Workbook:
            datemode = 0

            @staticmethod
            def sheets() -> list[Worksheet]:
                return [Worksheet()]

        xlrd = ModuleType("xlrd")
        xlrd.XL_CELL_NUMBER = 2  # type: ignore[attr-defined]
        xlrd.XL_CELL_DATE = 3  # type: ignore[attr-defined]
        xlrd.XLDateError = ValueError  # type: ignore[attr-defined]
        xlrd.xldate_as_datetime = lambda value, datemode: datetime(2023, 1, 15)  # type: ignore[attr-defined]
        xlrd.open_workbook = lambda *, file_contents: Workbook()  # type: ignore[attr-defined]
        monkeypatch.setitem(sys.modules, "xlrd", xlrd)

        sheets = CSVLoader._load_xls(b"legacy workbook")
        result = CSVLoader._format_excel_sheets(sheets, "report.xls", "xls")

        assert "Sheet: Legacy" in result.content
        assert "Row 1: name: South | revenue: 80 | opened: 2023-01-15" in result.content
        assert "80.0" not in result.content
        assert result.metadata["format"] == "xls"

    def test_xls_time_only_cell_omits_epoch_date(self):
        import xlrd

        cell = type("Cell", (), {"ctype": xlrd.XL_CELL_DATE, "value": 0.5})()

        assert CSVLoader._xls_cell_value(cell, 0) == "12:00:00"

    def test_xls_out_of_range_date_keeps_raw_value(self):
        import xlrd

        cell = type("Cell", (), {"ctype": xlrd.XL_CELL_DATE, "value": 1e20})()

        assert CSVLoader._xls_cell_value(cell, 0) == 1e20

    @patch("crewai_tools.security.safe_requests.safe_get_bounded")
    def test_load_xlsx_from_url(self, mock_get):
        buffer = BytesIO()
        workbook = Workbook()
        workbook.active.append(["name"])
        workbook.active.append(["Remote"])
        workbook.save(buffer)
        mock_get.return_value = (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "https://example.com/report.xlsx",
        )

        result = CSVLoader().load(
            SourceContent("https://example.com/report.xlsx?token=1")
        )

        assert "Row 1: name: Remote" in result.content
        assert mock_get.call_args.kwargs["max_bytes"] == _EXCEL_URL_MAX_BYTES
        assert "application/vnd.ms-excel" in mock_get.call_args.kwargs["headers"][
            "Accept"
        ]

    @patch("crewai_tools.security.safe_requests.safe_get_bounded")
    def test_xlsx_url_over_size_limit_is_refused(self, mock_get):
        mock_get.side_effect = ValueError("exceeds the 5242880 byte limit")

        with pytest.raises(ValueError, match="Error fetching content from URL"):
            CSVLoader().load(SourceContent("https://example.com/huge.xlsx"))
