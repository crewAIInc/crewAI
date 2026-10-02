import tempfile
from unittest.mock import Mock, patch

from docx import Document
from docx.document import Document as DocumentObject

from crewai_tools.rag.base_loader import LoaderResult
from crewai_tools.rag.loaders.docx_loader import DOCXLoader
from crewai_tools.rag.source_content import SourceContent
import pytest


def make_document(*paragraphs: str, tables: int = 0) -> DocumentObject:
    """Build a real DOCX document with the requested paragraphs and tables."""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for index in range(tables):
        document.add_table(rows=1, cols=1).cell(0, 0).text = f"Table {index + 1}"
    return document


class TestDOCXLoader:
    @patch("docx.Document")
    def test_load_docx_from_file(self, mock_docx_class):
        """Read nonempty paragraphs and preserve local document metadata."""
        mock_doc = make_document("First paragraph", "Second paragraph", "   ")
        mock_docx_class.return_value = mock_doc

        with tempfile.NamedTemporaryFile(suffix=".docx") as f:
            loader = DOCXLoader()
            result = loader.load(SourceContent(f.name))

            assert isinstance(result, LoaderResult)
            assert result.content == "First paragraph\nSecond paragraph"
            assert result.metadata == {"format": "docx", "paragraphs": 3, "tables": 0}
            assert result.source == f.name

    @patch("docx.Document")
    def test_load_docx_with_tables(self, mock_docx_class):
        """Count top-level tables in document metadata."""
        mock_doc = make_document("Document with table", tables=2)
        mock_docx_class.return_value = mock_doc

        with tempfile.NamedTemporaryFile(suffix=".docx") as f:
            loader = DOCXLoader()
            result = loader.load(SourceContent(f.name))

            assert result.metadata["tables"] == 2

    @patch("crewai_tools.security.safe_requests._raw_get")
    @patch("docx.Document")
    @patch("tempfile.NamedTemporaryFile")
    @patch("os.unlink")
    def test_load_docx_from_url(
        self, mock_unlink, mock_tempfile, mock_docx_class, mock_get
    ):
        """Download a DOCX file with the default document request headers."""
        mock_get.return_value = Mock(
            content=b"fake docx content", raise_for_status=Mock()
        )

        mock_temp = Mock(name="/tmp/temp_docx_file.docx")
        mock_temp.__enter__ = Mock(return_value=mock_temp)
        mock_temp.__exit__ = Mock(return_value=None)
        mock_tempfile.return_value = mock_temp

        mock_doc = make_document("Content from URL")
        mock_docx_class.return_value = mock_doc

        loader = DOCXLoader()
        result = loader.load(SourceContent("https://example.com/test.docx"))

        assert "Content from URL" in result.content
        assert result.source == "https://example.com/test.docx"

        headers = mock_get.call_args[1]["headers"]
        assert (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            in headers["Accept"]
        )
        assert "crewai-tools DOCXLoader" in headers["User-Agent"]

        mock_temp.write.assert_called_once_with(b"fake docx content")

    @patch("crewai_tools.security.safe_requests._raw_get")
    @patch("docx.Document")
    def test_load_docx_from_url_with_custom_headers(self, mock_docx_class, mock_get):
        """Forward custom headers when downloading DOCX content."""
        mock_get.return_value = Mock(
            content=b"fake docx content", raise_for_status=Mock()
        )
        mock_docx_class.return_value = make_document()

        loader = DOCXLoader()
        custom_headers = {"Authorization": "Bearer token"}

        with patch("tempfile.NamedTemporaryFile"), patch("os.unlink"):
            loader.load(
                SourceContent("https://example.com/test.docx"), headers=custom_headers
            )

        assert mock_get.call_args[1]["headers"] == custom_headers

    @patch("crewai_tools.security.safe_requests._raw_get")
    def test_load_docx_url_download_error(self, mock_get):
        """Report download failures with the source URL."""
        mock_get.side_effect = Exception("Network error")

        loader = DOCXLoader()
        with pytest.raises(ValueError, match="Error fetching content from URL"):
            loader.load(SourceContent("https://example.com/test.docx"))

    @patch("crewai_tools.security.safe_requests._raw_get")
    def test_load_docx_url_http_error(self, mock_get):
        """Reject unsuccessful HTTP responses before parsing DOCX data."""
        mock_get.return_value = Mock(
            raise_for_status=Mock(side_effect=Exception("404 Not Found"))
        )

        loader = DOCXLoader()
        with pytest.raises(ValueError, match="Error fetching content from URL"):
            loader.load(SourceContent("https://example.com/notfound.docx"))

    def test_load_docx_invalid_source(self):
        """Reject input that is neither an existing file nor a URL."""
        loader = DOCXLoader()
        with pytest.raises(ValueError, match="Source must be a valid file path or URL"):
            loader.load(SourceContent("not_a_file_or_url"))

    @patch("docx.Document")
    def test_load_docx_parsing_error(self, mock_docx_class):
        """Wrap document parsing failures with a loader-specific error."""
        mock_docx_class.side_effect = Exception("Invalid DOCX file")

        with tempfile.NamedTemporaryFile(suffix=".docx") as f:
            loader = DOCXLoader()
            with pytest.raises(ValueError, match="Error loading DOCX file"):
                loader.load(SourceContent(f.name))

    @patch("docx.Document")
    def test_load_docx_empty_document(self, mock_docx_class):
        """Return empty content and zero counts for a blank document."""
        mock_docx_class.return_value = make_document()

        with tempfile.NamedTemporaryFile(suffix=".docx") as f:
            loader = DOCXLoader()
            result = loader.load(SourceContent(f.name))

            assert result.content == ""
            assert result.metadata == {"paragraphs": 0, "tables": 0, "format": "docx"}

    @patch("docx.Document")
    def test_docx_doc_id_generation(self, mock_docx_class):
        """Generate a stable identifier for unchanged document content."""
        mock_docx_class.return_value = make_document("Consistent content")

        with tempfile.NamedTemporaryFile(suffix=".docx") as f:
            loader = DOCXLoader()
            source = SourceContent(f.name)
            assert loader.load(source).doc_id == loader.load(source).doc_id
