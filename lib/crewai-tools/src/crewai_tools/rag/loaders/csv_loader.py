from collections.abc import Iterable
import csv
from datetime import datetime, time
from io import BytesIO, StringIO
from typing import Any, Final
from urllib.parse import urlparse

from crewai_tools.rag.base_loader import BaseLoader, LoaderResult
from crewai_tools.rag.loaders.utils import load_from_url
from crewai_tools.rag.source_content import SourceContent


# Same ceiling URLReadTool uses for a remote workbook. A download limit does
# not bound how far a sheet expands once parsed.
_EXCEL_URL_MAX_BYTES: Final[int] = 5 * 1024 * 1024


class CSVLoader(BaseLoader):
    def load(self, source_content: SourceContent, **kwargs: Any) -> LoaderResult:  # type: ignore[override]
        source_ref = source_content.source_ref
        # urlparse treats '#' in a local name as a fragment, so "report#1.xlsx"
        # would miss the Excel branch. Only URLs should be parsed that way.
        source_path = (
            urlparse(source_ref).path if source_content.is_url() else source_ref
        )
        suffix = source_path.rsplit(".", 1)[-1].lower()

        if suffix in {"xls", "xlsx"}:
            content = self._load_excel_content(source_content, kwargs)
            sheets = (
                self._load_xlsx(content)
                if suffix == "xlsx"
                else self._load_xls(content)
            )
            return self._format_excel_sheets(sheets, source_ref, suffix)

        content_str = source_content.source
        if source_content.is_url():
            content_str = load_from_url(
                content_str,
                kwargs,
                accept_header="text/csv, application/csv, text/plain",
                loader_name="CSVLoader",
            )
        elif source_content.path_exists():
            content_str = self._load_from_file(content_str)

        return self._parse_csv(content_str, source_ref)

    @staticmethod
    def _load_excel_content(
        source_content: SourceContent, kwargs: dict[str, Any]
    ) -> bytes:
        if source_content.is_url():
            from crewai_tools.security.safe_requests import safe_get_bounded

            headers = kwargs.get(
                "headers",
                {
                    "Accept": "application/vnd.ms-excel, application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "User-Agent": "Mozilla/5.0 (compatible; crewai-tools CSVLoader)",
                },
            )
            try:
                body, _content_type, _final_url = safe_get_bounded(
                    source_content.source,
                    max_bytes=kwargs.get("max_bytes", _EXCEL_URL_MAX_BYTES),
                    headers=headers,
                    timeout=30,
                )
                return body
            except Exception as e:
                raise ValueError(
                    f"Error fetching content from URL {source_content.source}: {e!s}"
                ) from e

        with open(source_content.source, "rb") as file:
            return file.read()

    @staticmethod
    def _load_from_file(path: str) -> str:
        with open(path, encoding="utf-8") as file:
            return file.read()

    @staticmethod
    def _load_xlsx(content: bytes) -> list[tuple[str, Iterable[tuple[Any, ...]]]]:
        try:
            from openpyxl import load_workbook  # type: ignore[import-untyped]
        except ImportError as e:
            raise ImportError(
                "Reading .xlsx files requires openpyxl. Install with: uv add openpyxl"
            ) from e

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            return [
                (worksheet.title, list(worksheet.iter_rows(values_only=True)))
                for worksheet in workbook.worksheets
            ]
        finally:
            workbook.close()

    @staticmethod
    def _load_xls(content: bytes) -> list[tuple[str, Iterable[tuple[Any, ...]]]]:
        try:
            import xlrd  # type: ignore[import-untyped]
        except ImportError as e:
            raise ImportError(
                "Reading .xls files requires xlrd. Install with: uv add xlrd"
            ) from e

        workbook = xlrd.open_workbook(file_contents=content)
        return [
            (
                worksheet.name,
                [
                    tuple(
                        CSVLoader._xls_cell_value(cell, workbook.datemode)
                        for cell in worksheet.row(row)
                    )
                    for row in range(worksheet.nrows)
                ],
            )
            for worksheet in workbook.sheets()
        ]

    @staticmethod
    def _xls_cell_value(cell: Any, datemode: int) -> Any:
        """Return a searchable value for one xlrd cell.

        ``row_values()`` turns dates into Excel serial floats and whole numbers
        into ``80.0``. Dates become ISO text, and integral numbers stay ints.
        """
        import xlrd

        if cell.ctype == xlrd.XL_CELL_DATE:
            try:
                converted = xlrd.xldate_as_datetime(cell.value, datemode)
            except (xlrd.XLDateError, OverflowError):
                return cell.value
            if not isinstance(converted, datetime):
                return converted
            # A serial in [0, 1) is a time of day. xldate_as_datetime still
            # attaches the 1899/1904 epoch, which would be indexed as a date.
            if isinstance(cell.value, int | float) and 0 <= cell.value < 1:
                return converted.time().isoformat()
            if converted.time() == time.min:
                return converted.date().isoformat()
            return converted.isoformat(sep=" ")

        if cell.ctype == xlrd.XL_CELL_NUMBER and isinstance(cell.value, float):
            if cell.value.is_integer():
                return int(cell.value)
        return cell.value

    def _parse_csv(self, content: str, source_ref: str) -> LoaderResult:
        try:
            csv_reader = csv.DictReader(StringIO(content))

            text_parts = []
            headers = csv_reader.fieldnames

            if headers:
                text_parts.append("Headers: " + " | ".join(headers))
                text_parts.append("-" * 50)

                for row_num, row in enumerate(csv_reader, 1):
                    row_text = " | ".join([f"{k}: {v}" for k, v in row.items() if v])
                    text_parts.append(f"Row {row_num}: {row_text}")

            text = "\n".join(text_parts)

            metadata = {
                "format": "csv",
                "columns": headers,
                "rows": len(text_parts) - 2 if headers else 0,
            }

        except Exception as e:
            text = content
            metadata = {"format": "csv", "parse_error": str(e)}

        return LoaderResult(
            content=text,
            source=source_ref,
            metadata=metadata,
            doc_id=self.generate_doc_id(source_ref=source_ref, content=text),
        )

    @staticmethod
    def _format_excel_sheets(
        sheets: list[tuple[str, Iterable[tuple[Any, ...]]]],
        source_ref: str,
        suffix: str,
    ) -> LoaderResult:
        text_parts: list[str] = []
        sheet_metadata: list[dict[str, Any]] = []

        for sheet_name, rows in sheets:
            values = [
                ["" if value is None else str(value) for value in row] for row in rows
            ]
            if not values:
                sheet_metadata.append({"name": sheet_name, "columns": [], "rows": 0})
                continue

            headers = values[0]
            row_count = 0
            text_parts.extend(
                [f"Sheet: {sheet_name}", "Headers: " + " | ".join(headers), "-" * 50]
            )
            for row_num, row in enumerate(values[1:], 1):
                row_text = " | ".join(
                    f"{header}: {value}"
                    for header, value in zip(headers, row, strict=False)
                    if value
                )
                if row_text:
                    text_parts.append(f"Row {row_num}: {row_text}")
                    row_count += 1
            sheet_metadata.append(
                {"name": sheet_name, "columns": headers, "rows": row_count}
            )

        text = "\n".join(text_parts)
        return LoaderResult(
            content=text,
            source=source_ref,
            metadata={"format": suffix, "sheets": sheet_metadata},
            doc_id=CSVLoader.generate_doc_id(source_ref=source_ref, content=text),
        )
